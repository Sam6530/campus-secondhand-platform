from pathlib import Path
from hashlib import sha256
import re, json, subprocess, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
REPO = ROOT / 'campus-secondhand-platform'
ANNEX = OUT / '附件'
QA = OUT / '排版核对'
ANNEX.mkdir(exist_ok=True)
QA.mkdir(exist_ok=True)
SOURCE = REPO / 'documents/校园二手交易与物品流转平台需求规格说明书.md'
draft_path = OUT / 'D章节补充稿.md'
draft = draft_path.read_text(encoding='utf-8')
base = json.loads((OUT / '编写前仓库快照.json').read_text(encoding='utf-8'))

def local_link(path):
    import os
    return Path(os.path.relpath(path, ANNEX)).as_posix()

rows = [
('01','US06；UC01-a；账户；BR01、14','AC-D-01','SC-C-05仅有访问拒绝，不含注册'),
('02','US06；UC01-b；校园认证、处置记录；BR01、10、16','AC-D-01','未实现认证提交、退回和重审'),
('03','US02；UC04；物品；BR02、03；ST-B-01','AC-D-02','种子物品已审核，发布未验证'),
('04','US05-b；UC10-a；物品、处置记录；BR03、16；ST-B-02/03','AC-D-02','完整发布审核未验证'),
('05','US02；UC04；物品、订单快照；BR04；ST-B-05','AC-D-02','SC-C-21只测试内部资料变动与快照，不含修改重审'),
('06','US02；UC04；物品；BR05；ST-B-06/07','AC-D-02','无撤下、重发页面，未验证'),
('07','US01；UC02；物品、评价；BR03、11、13、14','AC-D-03','SC-C-21为公开字段与快照检查；关键词、分类筛选未验证'),
('08','US01、US04-a；UC03、06；收藏、会话；BR12、14、17','AC-D-04','SC-C-06/11/18涉及归属、旧单沟通和只读；收藏与新咨询未验证'),
('09','US03；UC05；订单、物品、快照；BR06、10、12、20；ST-B-08/13、S1/S2','AC-D-05','SC-C-01—05/12；仅顺序与故障模拟'),
('10','US04-a；UC06；订单、物品；BR07；ST-B-14/15/09','AC-D-06','SC-C-01/22；含接单、拒单及部分限制'),
('11','US04-b；UC07；订单、交付登记；BR09、10；ST-B-16/18、S3','AC-D-08','SC-C-07—09/13/14；仅先后顺序模拟'),
('12','US04-b/c；UC07、10-b；订单、物品；BR05、12、13、19；ST-B-09/11','AC-D-06/08/09','SC-C-08/10/15/22；不代表全部可售条件组合'),
('13','US04-a；UC06；订单、交付登记；BR08；ST-B-17','AC-D-07','SC-C-01/09/11；交付事实为模拟声明'),
('14','US04-a；UC06；订单、物品；BR08、14；ST-B-19/10、S4','AC-D-07','SC-C-01/06/18；无真实身份认证'),
('15','US04-a/c；UC06、10-b；订单、会话、争议；BR14、16、19','AC-D-05/07/09','SC-C-02/06/13/16；覆盖查询、提醒和部分权限'),
('16','US05-a；UC08；评价、订单；BR11、14','AC-D-10','完整评价资格、重复与公开治理未实现'),
('17','US05-a；UC09；举报、处置记录；BR15、16；ST-B-S6','AC-D-10','SC-C-19仅终态反馈、独立回执与事件关联模拟'),
('18','US05-b；UC10-b；举报、处置记录；BR15、16','AC-D-10','SC-C-19仅关联示例；举报证据核查及完整审核未验证'),
('19','US05-b；UC11；账户、处置记录；BR12、16','AC-D-11','SC-C-10—12/20涉及限制、恢复或权限边界；无真实治理服务'),
('20','US05-b；UC11；账户、物品、订单；BR12、13；ST-B-S1','AC-D-05/11','SC-C-02/10—12/22；旧单必要权利与新业务限制模拟'),
('21','US05-c；UC12；订单、快照与统计计算；BR18；ST-B-19/10','AC-D-12','SC-C-01/15/18/19为完成去重与金额；完整报表和周期未验证'),
('22','US05-b；UC04—12；处置记录、订单历史；BR16','AC-D-02/05/12','程序逐步记录与日志可查；真实持久化和完整审计未验证'),
('23','US06；UC01-c/d；账户、会话；BR01、14','AC-D-01','角色切换不是真实登录、退出或会话失效验证'),
('24','US04-c；UC10-b；履约争议、订单、物品；BR15、19；ST-B-20/21/11','AC-D-09','SC-C-13—15及UI-C-08涉及恢复、交付冲突和终止模拟'),
]
trace = '# E-D-02 需求追踪明细\n\n'
trace += '版本：D-trace-v0.1。日期：2026年10月22日。输入提交：'+base['git']['head']+'。本表为静态映射，验收标准来自本次D章节草稿；A/B/C尚未对该标准完成交叉核对。C的SC/PV统一按C-prototype-v1.1对应记录理解，映射不表示整项需求已验证。\n\n'
trace += '| 需求 | 故事、用例、对象与规则／状态 | 验收标准 | 已有证据及范围限制 |\n| --- | --- | --- | --- |\n'
for fr, refs, ac, evidence in rows:
    trace += f'| FR{fr} | {refs} | {ac} | {evidence} |\n'
trace += '\nST-B-S1—S6是状态保持或拒绝记录，不是成功转换；无业务生命周期变化的用例不强行添加ST号。SC-C-21的内部资料修改是测试夹具，不代表产品具有编辑页面。\n\n'
trace += '资料来源SR01的十类功能对应以E-A-03“课程功能与需求对应”为入口；各BR确认状态引用该原件，不由本表升级。每条FR的原文仍以需求正文2.1为准。\n\n'
trace += '核查方法：逐条读取FR01—24及US/UC关系，对照3.2—3.5对象职责、第4章、第5章、C场景名称与范围，分配AC-D-01—12。仅检查有对应项、编号有效与证据边界，未以本表证明系统完整性或取代后续原型操作。\n'
(ANNEX/'需求追踪明细.md').write_text(trace, encoding='utf-8')

evidence_paths = [
'documents/校园二手交易与物品流转平台需求规格说明书.md',
'documents/需求分析质量验证报告.md',
'initialization/skills/req-d-verify/SKILL.md',
'initialization/skills/共同大纲.md','initialization/skills/协作协议.md',
'initialization/docs/提交要求对照.md','initialization/docs/文档格式规范.md',
'attachments/A/依据/实践项目说明(1).pdf',
'attachments/A/依据/学生需求分析与AI智能体使用验证报告.md',
'attachments/A/依据与规则核验记录.md','attachments/A/PR回复.md',
'attachments/B/模型与规则依据核验记录.md','attachments/B/模型说明与状态转换核对.md',
'attachments/C/原型范围与运行说明.md','attachments/C/场景数据与走查脚本.md',
'attachments/C/问题清单与复验.md','attachments/C/原型验证记录.md',
'attachments/C/修订说明与审核交接.md','attachments/总证据索引.md','attachments/交付清单.md',
]
record = '# E-D-01 编写依据与核对记录\n\n'
record += '任务号：AI-D-CH-01；版本：D-chapters-v0.1；日期：2026年10月22日（Asia/Shanghai）。实际操作者：D侧AI助手；人工授权者：本聊天成员D，实名未提供。本稿尚无D逐条审阅结论。\n\n'
record += '## 本次授权及位置\n\n用户原话：\n\n> 嗯，有关于补全章节的部分，你去补充罢，注意，目前所做的一切都不要交仓库里，我知道你有这个权限，但你不要这么做，ok？\n\n'
record += '本次只补7.1、7.2、7.4与第8章，产物保存在本目录；源仓库只读，不更新远端、不切换分支、不编辑源文件、不提交或推送。本地输入分支为'+base['git']['branch']+'，完整提交号为`'+base['git']['head']+'`。本次没有重新核实GitHub main是否已合并C，因此不将本地C分支快照称为最新正式main。\n\n'
record += '## 输入文件与指纹\n\n以下指纹按本地文件原始字节计算，换行格式可能与其他系统的LF指纹不同。\n\n| 材料 | SHA256 |\n| --- | --- |\n'
for rel in evidence_paths:
    p = REPO/rel
    assert p.exists(), p
    record += f'| [{rel}]({local_link(p)}) | `{sha256(p.read_bytes()).hexdigest()}` |\n'
record += '\n## 已执行的核对与可用范围\n\n'
record += '1. 读取课程PDF第2页选题原文，核对十类功能与A的FR映射；读取原始提交要求及共同大纲v3.3，按D职责起草。没有开展访谈或取得新的业务确认。\n'
record += '2. 建立AC-D-01—12及FR01—24逐条映射，保留未实现、未执行与待确认范围；未改变BR或ST的权威定义。D对新标准的人工审阅及A/B/C交叉核对未开展。\n'
record += '3. 引用本聊天本日前一轮真实执行的C复核：D侧AI在仓库外副本运行verify.cjs，22组73断言通过；再运行walkthrough-check.cjs，22组文档调用重放通过。此处只是将已有执行纳入可追踪记录，本次章节起草不冒充再次运行。\n'
for filename in ['程序验证结果.json','脚本复现检查.json']:
    p=ROOT/'tmp/c_v2_review_20261022/verification/revision-20261022'/filename
    assert p.exists(), p
    data=json.loads(p.read_text(encoding='utf-8'))
    record += f'   - [{filename}]({local_link(p)})：执行时间`{data["date"]}`，{data["passed"]}/{data["total"]}通过，SHA256 `{sha256(p.read_bytes()).hexdigest()}`。\n'
record += '4. C提交的8组页面结果在前一轮经源文件指纹、检查脚本及未登记终止截图核对；D侧未重新执行浏览器。真实后台并发、认证、持久化、通知与性能未验证，学生人工走查未记录。\n'
record += '5. 正文7.3保留C的原文，仍是2026年10月21日首轮摘要；7.1和8.1已指明10月22日增量。后续正式汇编应由C统一更新摘要，不在本次越出补章范围。\n\n'
record += '## 人工取舍的真实依据\n\n'
record += '- 成员D在本聊天说：“嗯，我同意了，就把这几项加进来吧，但不算错误，只能算是建议”。原发言逐条日期未记录，不推定日期；该意见只支持建议分类，不代表已作逐条业务确认。\n'
record += '- B轮次记录保存了成员D的“我认为问题不大，我打算在我这一轮再作调整，现在先给他过了”，日期依据为2026年10月21日记录。\n'
p=ROOT/'D核验/2026-10-21_B轮次决定.md'
record += f'  原件：[B轮次决定]({local_link(p)})，SHA256 `{sha256(p.read_bytes()).hexdigest()}`。该记录列明保留项，并未将本轮通过等同问题修复。\n'
record += '- 两项C业务口径的具体讨论沿用E-C-06附件；按成员D当前意见，D正文不展开为新增缺陷，不将C演示选择直接写成正式政策。\n\n'
record += '## 本次实际动作与限制\n\n完成章节起草、24条追踪映射、输入版本登记及仓库外预览制作。未开展新的多agent审查，不虚构两个独立子agent已运行；先前六角色审查的完整过程本次未重建或补造。专门AI使用报告、完整独立对抗审查与正式验收仍属于后续工作。\n\n'
record += '本稿表号使用7-D/8-D前缀，以免改写C既有表7-1；正式合并时由C统一按出现顺序编号。标准按当前正文派生，不能单凭程序能运行判断业务正确，需D实际核查、A/B/C交叉反馈后修订。\n\n'
record += '仓库不变核对：编写前快照为本目录《编写前仓库快照.json》；最终核对结果写入《完成核对.json》，比较全部非.git文件原始字节、HEAD、分支、refs及工作区状态。\n'
(ANNEX/'编写依据与核对记录.md').write_text(record, encoding='utf-8')

original = SOURCE.read_text(encoding='utf-8-sig')
integrated = original
titles = ['7.1 核心需求验收场景与标准','7.2 资料来源核对与规则确认','7.4 覆盖、追踪与一致性检查']
for title in titles:
    pattern = re.compile(r'^### '+re.escape(title)+r'\n.*?(?=^### |^## |\Z)',re.M|re.S)
    match = pattern.search(draft)
    assert match, title
    integrated,n=pattern.subn(lambda _:match.group(0).rstrip()+'\n\n',integrated,count=1)
    assert n==1
chapter8 = draft.split('## 8 基线与人机协作\n',1)[1]
integrated,n=re.subn(r'^## 8 基线与人机协作\n.*\Z','## 8 基线与人机协作\n'+chapter8,integrated,count=1,flags=re.M|re.S)
assert n==1
def protected(text):
    for title in titles:
        text=re.sub(r'^### '+re.escape(title)+r'\n.*?(?=^### |^## |\Z)', 'SECTION\n',text,flags=re.M|re.S)
    text=re.sub(r'^## 8 基线与人机协作\n.*\Z','CH8\n',text,flags=re.M|re.S)
    return text.strip()
assert protected(original)==protected(integrated),'A/B/C content changed'
# Rebase existing figure links for an external preview without copying source evidence.
def rebase(m):
    label,target=m.groups()
    if re.match(r'^(https?:|#|/|[A-Za-z]:)',target):return m.group(0)
    p=(SOURCE.parent/target).resolve()
    assert p.exists(), p
    import os
    return f'{label}({Path(os.path.relpath(p,OUT)).as_posix()})'
integrated=re.sub(r'(!?\[[^\]]*\])\(([^)]+)\)',rebase,integrated)
integrated=integrated.replace('D补充材料E-D-01','D补充材料[E-D-01](附件/编写依据与核对记录.md)').replace('D补充材料E-D-02','D补充材料[E-D-02](附件/需求追踪明细.md)')
(OUT/'需求规格说明书_补入D章节预览.md').write_text(integrated,encoding='utf-8')

def font(style, size, bold=False):
    style.font.name='楷体'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'楷体')

doc=Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
sec.header_distance=sec.footer_distance=Cm(0.8)
normal=doc.styles['Normal']
font(normal,12)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.widow_control=True
for name,size in [('Title',20),('Heading 1',16),('Heading 2',14)]:
    font(doc.styles[name],size,True)
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles[name].paragraph_format.keep_with_next=True
font(doc.styles['Caption'],11,True)
doc.styles['Caption'].paragraph_format.keep_with_next=True
doc.styles['Caption'].paragraph_format.space_before=Pt(6)
doc.styles['Caption'].paragraph_format.space_after=Pt(4)
doc.core_properties.title='校园二手交易与物品流转平台：D章节补充稿'
doc.core_properties.author='成员D（AI辅助起草）'
doc.core_properties.subject='7.1、7.2、7.4及第8章'
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
expected=[]
lines=draft.splitlines();i=0;tables=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[i+1].strip()):
        data=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=lines[i].strip()
            if not re.match(r'^\|[\s:|\-]+\|$',row):data.append([x.strip() for x in row.split('|')[1:-1]])
            i+=1
        assert all(len(x)==len(data[0]) for x in data)
        table=doc.add_table(rows=1,cols=len(data[0]))
        table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        widths=([3.0,5.0,6.1,2.9] if len(data[0])==4 else [4.2,6.7,6.1])
        # A fixed 17cm body width. Slightly wider result column in the acceptance table.
        assert abs(sum(widths)-17)<0.01
        for col,width in zip(table.columns,widths):col.width=Cm(width)
        for ri,row in enumerate(data):
            cells=table.rows[0].cells if ri==0 else table.add_row().cells
            for ci,txt in enumerate(row):
                cells[ci].width=Cm(widths[ci]);cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cells[ci].paragraphs[0];p.paragraph_format.line_spacing=1.5
                p.paragraph_format.keep_with_next=(ri==0)
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                run=p.add_run(txt);run.font.size=Pt(10.5);run.bold=(ri==0)
                expected.append(txt)
            props=table.rows[ri]._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if ri==0:
                h=OxmlElement('w:tblHeader');h.set(qn('w:val'),'true');props.append(h)
                for c in cells:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F3F5');c._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        tables+=1
        continue
    h=re.match(r'^(#{1,3})\s+(.*)$',line)
    if h:
        level=len(h.group(1));txt=h.group(2)
        p=doc.add_paragraph(txt,style={1:'Title',2:'Heading 1',3:'Heading 2'}[level])
        if level==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    else:
        txt=line
        p=doc.add_paragraph(txt,style='Caption' if re.match(r'^表[78]-D\d+ ',txt) else 'Normal')
        if not txt.startswith('表') and not txt.startswith('版本：'):
            p.paragraph_format.first_line_indent=Pt(24)
    expected.append(txt)
    i+=1
docx=OUT/'D章节补充稿.docx';doc.save(docx)
manifest={'source_commit':base['git']['head'],'source_md_sha256':sha256(SOURCE.read_bytes()).hexdigest(),'draft_sha256':sha256(draft_path.read_bytes()).hexdigest(),'protected_original_content_unchanged':True,'FR_count':len(rows),'AC_count':len(set(re.findall(r'AC-D-\d+',draft))),'table_count':tables,'expected_fragments':expected}
(QA/'构建记录.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k!='expected_fragments'},ensure_ascii=False,indent=2))
print('Draft chars',len(draft),'DOCX',docx)
